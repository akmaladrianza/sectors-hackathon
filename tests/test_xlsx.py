"""Tests for the formula-driven DCF xlsx export (pure, in-memory workbook).

Run:  python -m tests.test_xlsx
"""

from __future__ import annotations

import io
from decimal import Decimal

import pandas as pd

from engine.xlsx_export import write_dcf_sheet, _col


def _build_and_open(**kw) -> "openpyxl.Workbook":
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        base = dict(
            revenue=Decimal("1000"),
            g=Decimal("0.08"),
            m=Decimal("0.40"),
            r=Decimal("0.12"),
            tvg=Decimal("0.02"),
            years=5,
        )
        base.update(kw)
        write_dcf_sheet(writer, "DCF T", **base)
    buf.seek(0)
    import openpyxl

    return openpyxl.load_workbook(buf)


def _assumptions(ws):
    return {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
        for r in range(2, 20)
        if ws.cell(row=r, column=1).value
    }


def test_nwc_change_margin_is_not_re_divided() -> None:
    """ΔNWC must land in the sheet verbatim (it is already a margin, not a dollar amount).

    Regression for a units bug: ``nwc_change_margin_from_days`` returns ΔNWC as a
    *fraction of revenue* (e.g. -0.0010147…), and the xlsx writer previously divided it
    by revenue *again*, producing a near-zero garbage value. The cell must equal the
    input margin exactly.
    """
    margin = Decimal("-0.001014713343480466768138001015")
    wb = _build_and_open(nwc_change_margin=margin)
    ws = wb["DCF T"]
    vals = _assumptions(ws)
    assert "ΔNWC (of revenue)" in vals, vals
    actual = vals["ΔNWC (of revenue)"]
    assert actual is not None
    assert abs(actual - float(margin)) < 1e-12, actual
    print(f"PASS ΔNWC cell equals input margin ({actual}), not a stale /revenue double-division\n")


def test_net_debt_and_shares_are_visible_assumptions() -> None:
    """Net debt + shares must be their own labelled cells, referenced by the FCFE bridge."""
    wb = _build_and_open(
        shares_outstanding=Decimal("99061024659"),
        net_debt=Decimal("16546000000000"),
    )
    ws = wb["DCF T"]
    vals = _assumptions(ws)
    assert "Net debt" in vals and vals["Net debt"] == 16546000000000.0, vals
    assert "Shares outstanding" in vals and vals["Shares outstanding"] == 99061024659.0, vals

    # The FCFE bridge must reference the cells, not embed literals.
    formulas = [ws.cell(row=r, column=12).value for r in range(20, 30)]
    eq_formula = next((f for f in formulas if isinstance(f, str) and "-$B$" in f), None)
    assert eq_formula is not None and "-16546000000000.0" not in eq_formula, formulas
    print("PASS net debt + shares are visible, referenced (not inlined) in FCFE bridge\n")


def test_ar_inv_ap_days_rows_present() -> None:
    wb = _build_and_open(ar_days=45.0, inv_days=10.0, ap_days=60.0)
    ws = wb["DCF T"]
    vals = _assumptions(ws)
    assert vals["AR days"] == 45.0
    assert vals["Inventory days"] == 10.0
    assert vals["AP days"] == 60.0
    print("PASS AR/Inventory/AP days surfaced as explicit assumption rows\n")


def main() -> None:
    test_nwc_change_margin_is_not_re_divided()
    test_net_debt_and_shares_are_visible_assumptions()
    test_ar_inv_ap_days_rows_present()
    print("ALL XLSX TESTS PASSED")


if __name__ == "__main__":
    main()
